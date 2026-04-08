"""Unit tests for Excel export logic."""
import pytest
from openpyxl import Workbook


def test_write_text_clip():
    from excel_export.clip_writer import write_text_clip
    wb = Workbook()
    ws = wb.active
    write_text_clip(ws, "B5", "Invoice Total")
    assert ws["B5"].value == "Invoice Total"


def test_write_text_clip_with_sheet_prefix():
    from excel_export.clip_writer import write_text_clip
    wb = Workbook()
    ws = wb.active
    write_text_clip(ws, "Sheet1!C3", "Test Value")
    assert ws["C3"].value == "Test Value"


def test_write_table_clip():
    from excel_export.clip_writer import write_table_clip
    wb = Workbook()
    ws = wb.active
    data = [["Item", "Qty", "Price"], ["Widget", "10", "100"], ["Gadget", "5", "250"]]
    write_table_clip(ws, "B2:D4", data)
    assert ws["B2"].value == "Item"
    assert ws["C2"].value == "Qty"
    assert ws["B3"].value == "Widget"
    assert ws["D4"].value == "250"


def test_write_calc_clip():
    from excel_export.formula_writer import write_calc_clip
    wb = Workbook()
    ws = wb.active
    write_calc_clip(ws, "B10", "=SUM(B2,B3,B4)")
    assert ws["B10"].value == "=SUM(B2,B3,B4)"


def test_validation_style_validated():
    from excel_export.validation_styler import apply_validation_style, VALIDATED_FILL
    wb = Workbook()
    ws = wb.active
    cell = ws["A1"]
    apply_validation_style(cell, "validated")
    assert cell.fill.fgColor.rgb == VALIDATED_FILL.fgColor.rgb


def test_validation_style_exception():
    from excel_export.validation_styler import apply_validation_style, EXCEPTION_FILL
    wb = Workbook()
    ws = wb.active
    cell = ws["A1"]
    apply_validation_style(cell, "exception")
    assert cell.fill.fgColor.rgb == EXCEPTION_FILL.fgColor.rgb


def test_table_merger_primary_wins():
    from table_extraction.table_merger import reconcile_tables
    primary = [[["A", "B"], ["1", "2"]]]
    fallback = [[["C", "D"]]]
    assert reconcile_tables(primary, fallback) == primary


def test_table_merger_uses_fallback_when_primary_empty():
    from table_extraction.table_merger import reconcile_tables
    fallback = [[["C", "D"]]]
    assert reconcile_tables([], fallback) == fallback


def test_normalize_table():
    from table_extraction.table_merger import normalize_table
    raw = [["  hello  ", " world "], ["a"]]
    normalized = normalize_table(raw)
    assert normalized[0][0] == "hello"
    assert normalized[0][1] == "world"
    assert normalized[1][0] == "a"
    assert normalized[1][1] == ""  # padded
