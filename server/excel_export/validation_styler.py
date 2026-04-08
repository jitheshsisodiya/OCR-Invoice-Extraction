"""Apply validation color fills to Excel cells."""
from openpyxl.styles import PatternFill

VALIDATED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")   # green
EXCEPTION_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")   # red
PENDING_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")     # yellow


def apply_validation_style(cell, status: str) -> None:
    if status == "validated":
        cell.fill = VALIDATED_FILL
    elif status == "exception":
        cell.fill = EXCEPTION_FILL
    else:
        cell.fill = PENDING_FILL
