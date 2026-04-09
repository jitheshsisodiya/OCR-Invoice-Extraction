"""Create the 'Source Map' sheet for audit traceability."""
from __future__ import annotations
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_source_map_sheet(wb: Workbook, clips: list) -> None:
    """
    Add a 'Source Map' sheet listing every clip → source document + page + bbox.
    clips: list of ClipMapping ORM objects.
    """
    ws = wb.create_sheet("Source Map")
    ws.sheet_state = "visible"  # visible but separate tab

    # Header
    headers = ["Excel Cell/Range", "Clip Type", "Document", "Page", "BBox (x1,y1,x2,y2)", "Confidence", "Engine", "Validation"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, clip in enumerate(clips, 2):
        cell_ref = clip.excel_range or clip.excel_cell or ""
        bbox_str = ""
        if clip.bbox:
            try:
                # bbox is a JSON column (dict), but guard against legacy string values
                b = clip.bbox if isinstance(clip.bbox, dict) else json.loads(clip.bbox)
                bbox_str = f"{b.get('x1',0):.3f},{b.get('y1',0):.3f},{b.get('x2',0):.3f},{b.get('y2',0):.3f}"
            except Exception:
                pass

        validation_status = ""
        if clip.validation:
            validation_status = clip.validation.status

        # Get document filename from session → document relationship
        doc_name = ""
        if hasattr(clip, "session") and clip.session and hasattr(clip.session, "document"):
            doc_name = clip.session.document.filename if clip.session.document else ""

        ws.cell(row=row_idx, column=1, value=cell_ref)
        ws.cell(row=row_idx, column=2, value=clip.clip_type)
        ws.cell(row=row_idx, column=3, value=doc_name)
        ws.cell(row=row_idx, column=4, value=clip.page_number)
        ws.cell(row=row_idx, column=5, value=bbox_str)
        ws.cell(row=row_idx, column=6, value=round(clip.confidence, 3))
        ws.cell(row=row_idx, column=7, value=clip.ocr_engine)
        ws.cell(row=row_idx, column=8, value=validation_status)

    # Auto-width
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
