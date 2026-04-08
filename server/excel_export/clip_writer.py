"""Write text and table clips to Excel worksheets."""
from __future__ import annotations
import json
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def write_text_clip(ws: Worksheet, excel_cell: str, value: str) -> None:
    """Write a single text value to the given cell address (e.g. 'B5')."""
    # Strip sheet prefix if present (e.g. "Sheet1!B5" → "B5")
    cell_ref = excel_cell.split("!")[-1] if "!" in excel_cell else excel_cell
    ws[cell_ref] = value


def write_table_clip(ws: Worksheet, excel_range: str, table_data: list[list[str]]) -> None:
    """Write a 2D table starting at the top-left cell of excel_range."""
    start_cell = excel_range.split("!")[-1].split(":")[0] if excel_range else "A1"
    # Parse start cell row/col
    from openpyxl.utils import column_index_from_string, get_column_letter
    import re
    match = re.match(r"([A-Z]+)(\d+)", start_cell.upper())
    if not match:
        return
    start_col = column_index_from_string(match.group(1))
    start_row = int(match.group(2))

    for r_idx, row in enumerate(table_data):
        for c_idx, cell_val in enumerate(row):
            col_letter = get_column_letter(start_col + c_idx)
            ws[f"{col_letter}{start_row + r_idx}"] = str(cell_val)
